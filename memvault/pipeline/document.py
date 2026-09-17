"""文档解析:PDF / Word / Excel / PPT / 纯文本 → 语义块入库。

轻量实现(不依赖 unstructured):PyMuPDF + python-docx + openpyxl + python-pptx。
PDF 内嵌图片(超过阈值尺寸)抽取为图像块,其余格式先做文本层。
"""
import logging
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".log"}
DOC_EXT = {".pdf", ".docx", ".doc", ".xlsx", ".xlsm", ".xls",
           ".pptx", ".ppt"} | TEXT_EXT

_MAX_CHUNK = 1200        # 单块字符上限
_MAX_CHUNKS = 200        # 单文档块数上限
_MAX_PDF_PAGES = 80
_MAX_PDF_IMAGES = 10
_MIN_IMAGE_BYTES = 30_000  # 小于 30KB 的图视为图标/水印跳过


def _chunk_text(text: str) -> list[str]:
    """按段落聚合切块,避免块过小或超长。"""
    chunks, buf = [], ""
    for para in (text or "").split("\n"):
        para = para.strip()
        if not para:
            continue
        if len(buf) + len(para) + 1 > _MAX_CHUNK and buf:
            chunks.append(buf)
            buf = ""
        buf = (buf + "\n" + para) if buf else para
    if buf:
        chunks.append(buf)
    return chunks[:_MAX_CHUNKS]


def _parse_pdf(path: Path, images_dir: Path):
    try:
        import pymupdf as fitz
    except ImportError:  # 兼容旧版包名
        import fitz

    texts, images = [], []
    with fitz.open(str(path)) as doc:
        for pno, page in enumerate(doc):
            if pno >= _MAX_PDF_PAGES:
                break
            t = page.get_text().strip()
            if t:
                texts.append(f"[第{pno + 1}页]\n{t}")
        all_images = [i for page in doc for i in page.get_images(full=True)]
        for img_no, img in enumerate(all_images):
            if len(images) >= _MAX_PDF_IMAGES:
                break
            try:
                xref = img[0]
                base = doc.extract_image(xref)
                if not base or len(base.get("image", b"")) < _MIN_IMAGE_BYTES:
                    continue
                images_dir.mkdir(parents=True, exist_ok=True)
                fp = images_dir / f"pdf_img_{img_no:02d}.{base['ext']}"
                fp.write_bytes(base["image"])
                images.append(str(fp))
            except Exception as e:  # noqa: BLE001 — 单图失败不影响解析
                logger.warning("PDF 图片抽取失败: %s", e)
    return "\n".join(texts), images


def _parse_docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for tbl_no, table in enumerate(doc.tables):
        rows = [" | ".join(c.text.strip() for c in row.cells)
                for row in table.rows]
        parts.append(f"[表格{tbl_no + 1}]\n" + "\n".join(rows))
    return "\n".join(parts)


def _parse_excel(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        lines = [f"[工作表:{ws.title}]"]
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 200:
                lines.append("...(截断)")
                break
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
        parts.append("\n".join(lines))
    wb.close()
    return "\n".join(parts)


def _parse_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides):
        lines = [f"[第{i + 1}页]"]
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        if len(lines) > 1:
            parts.append("\n".join(lines))
    return "\n".join(parts)


def parse_document(path, memory, cfg: dict, domain: str = "general") -> int:
    """解析文档并入库,返回 item_id。"""
    from memvault.config import media_dir

    path = Path(path)
    ext = path.suffix.lower()
    if ext not in DOC_EXT:
        raise ValueError(f"不支持的文档格式: {ext}")

    item_id = memory.add_item(
        domain=domain, type_="doc",
        title=path.stem,
        source_type="file", source_ref=str(path.resolve()),
    )
    images_dir = media_dir(cfg) / f"item_{item_id}" / "images"

    if ext == ".pdf":
        text, images = _parse_pdf(path, images_dir)
    elif ext in (".docx", ".doc"):
        text, images = _parse_docx(path), []
    elif ext in (".xlsx", ".xlsm", ".xls"):
        text, images = _parse_excel(path), []
    elif ext in (".pptx", ".ppt"):
        text, images = _parse_pptx(path), []
    else:  # 纯文本类
        text = path.read_text(encoding="utf-8", errors="ignore")
        images = []

    chunks = _chunk_text(text)
    for seq, chunk in enumerate(chunks):
        memory.add_text_chunk(item_id, chunk, seq=seq)
    for img_seq, img in enumerate(images):
        memory.add_image_chunk(item_id, img, seq=500 + img_seq)

    memory.db.update_item_media(
        item_id, content_text=(text[:4000] or "(未提取到文本)"),
        media_paths=images)
    logger.info("文档入库 item=%s (%s): %d 块 + %d 图",
                item_id, ext, len(chunks), len(images))
    return item_id
