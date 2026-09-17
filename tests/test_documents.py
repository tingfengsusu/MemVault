"""文档解析测试:现场生成 docx/xlsx/pptx/pdf/txt,验证解析入库与检索。"""
import pytest


@pytest.fixture()
def doc_paths(tmp_path):
    return tmp_path / "docs"


def test_parse_docx(memory, cfg, doc_paths):
    import docx

    doc_paths.mkdir(exist_ok=True)
    p = doc_paths / "训练计划.docx"
    d = docx.Document()
    d.add_heading("八周增肌计划", 0)
    d.add_paragraph("周一练胸:平板卧推 5x5,上斜哑铃 4x10。")
    d.add_paragraph("周三练背:引体向上 4x6,杠铃划船 4x10。")
    table = d.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "动作"
    table.rows[0].cells[1].text = "组数"
    table.rows[1].cells[0].text = "深蹲"
    table.rows[1].cells[1].text = "5x5"
    d.save(str(p))

    from memvault.pipeline.document import parse_document

    item_id = parse_document(p, memory, cfg, domain="fitness")
    item = memory.get_item(item_id)
    assert item["type"] == "doc"
    all_text = "\n".join(c["content"] for c in item["chunks"])
    assert "平板卧推" in all_text and "[表格1]" in all_text
    assert memory.search("引体向上 4x6")  # 可检索


def test_parse_excel(memory, cfg, doc_paths):
    import openpyxl

    doc_paths.mkdir(exist_ok=True)
    p = doc_paths / "开支记录.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "1月"
    ws.append(["日期", "项目", "金额"])
    ws.append(["1月3日", "蛋白粉", 299])
    ws.append(["1月8日", "跑鞋", 899])
    wb.save(str(p))

    from memvault.pipeline.document import parse_document

    item_id = parse_document(p, memory, cfg, domain="shopping")
    item = memory.get_item(item_id)
    all_text = "\n".join(c["content"] for c in item["chunks"])
    assert "[工作表:1月]" in all_text and "蛋白粉" in all_text


def test_parse_pptx(memory, cfg, doc_paths):
    from pptx import Presentation

    doc_paths.mkdir(exist_ok=True)
    p = doc_paths / "季度汇报.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "第三季度总结"
    slide.placeholders[1].text = "营收增长 20%,重点是新客转化。"
    prs.save(str(p))

    from memvault.pipeline.document import parse_document

    item_id = parse_document(p, memory, cfg, domain="general")
    all_text = "\n".join(c["content"] for c in memory.get_item(item_id)["chunks"])
    assert "[第1页]" in all_text and "营收增长 20%" in all_text


def test_parse_pdf(memory, cfg, doc_paths):
    import pymupdf

    doc_paths.mkdir(exist_ok=True)
    p = doc_paths / "产品说明.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "MemVault 使用说明", fontsize=16,
                     fontname="china-s")  # 内置中文字体
    page.insert_text((72, 130), "支持视频、网页与文档的本地记忆管理。",
                     fontsize=11, fontname="china-s")
    doc.save(str(p))
    doc.close()

    from memvault.pipeline.document import parse_document

    item_id = parse_document(p, memory, cfg, domain="general")
    all_text = "\n".join(c["content"] for c in memory.get_item(item_id)["chunks"])
    assert "[第1页]" in all_text and "本地记忆管理" in all_text


def test_ingest_file_routes_document(memory, cfg, tmp_path):
    from memvault.pipeline.files import ingest_file

    p = tmp_path / "笔记.md"
    p.write_text("# 周记\n本周把快照备份脚本补完了。", encoding="utf-8")
    item_id = ingest_file({"path": str(p)}, memory, cfg)
    item = memory.get_item(item_id)
    assert item["type"] == "doc"
    assert "快照备份脚本" in (item["content_text"] or "")


def test_long_text_chunking(memory, cfg, tmp_path):
    from memvault.pipeline.document import _chunk_text

    text = "\n".join(f"第{i}段:" + "内容" * 300 for i in range(10))
    chunks = _chunk_text(text)
    assert len(chunks) >= 5  # 长文被切成多块
    assert all(len(c) <= 1400 for c in chunks)


def test_parse_epub(memory, cfg, doc_paths):
    import zipfile

    doc_paths.mkdir(exist_ok=True)
    p = doc_paths / "小书.epub"
    with zipfile.ZipFile(p, "w") as z:
        z.writestr("mimetype", "application/epub+zip")
        z.writestr("META-INF/container.xml",
            '<?xml version="1.0"?><container version="1.0" '
            'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
            '<rootfiles><rootfile full-path="content.opf" '
            'media-type="application/oebps-package+xml"/></rootfiles></container>')
        z.writestr("content.opf",
            '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" '
            'version="3.0"><metadata '
            'xmlns:dc="http://purl.org/dc/elements/1.1/">'
            '<dc:title>习惯的力量</dc:title></metadata>'
            '<manifest><item id="c1" href="ch1.xhtml" '
            'media-type="application/xhtml+xml"/></manifest>'
            '<spine><itemref idref="c1"/></spine></package>')
        z.writestr("ch1.xhtml",
            "<html><body><h1>第一章</h1>"
            "<p>习惯由提示、惯常行为和奖赏三部分组成。</p></body></html>")

    from memvault.pipeline.document import parse_document

    item_id = parse_document(p, memory, cfg, domain="reading")
    all_text = "\n".join(c["content"] for c in memory.get_item(item_id)["chunks"])
    assert "[书名]习惯的力量" in all_text
    assert "惯常行为" in all_text
